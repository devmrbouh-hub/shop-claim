# -*- coding: utf-8 -*-
"""Import offers-template.xlsx → catalog/offers.yaml + vehicle_profiles.yaml.

Run from this directory:
  python import-offers-from-xlsx.py

Requires: openpyxl, pyyaml
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
CATALOG_DIR = SCRIPT_DIR.parent
DEFAULT_INPUT = SCRIPT_DIR / "offers-template.xlsx"
DEFAULT_MARKET_DIR = (
    SCRIPT_DIR.parents[3] / "Instance_1" / "ExpansionMod" / "Market"
)

SKIP_OFFER_IDS = {12345, 12346, 12350, 12351}

VEHICLE_OVERRIDES: dict[int, str] = {
    5010: "OffroadHatchback",
}

DEFAULT_SPAWN = {
    "distance_m": 6,
    "max_slope_deg": 15,
    "check_radius_m": 2.5,
}

KNOWN_CLASS_ALIASES: dict[str, str] = {
    "sparkplug": "SparkPlug",
    "headlighth7": "HeadlightH7",
    "carradiator": "CarRadiator",
    "carbattery": "CarBattery",
    "truckbattery": "TruckBattery",
}


def _parse_offer_id(raw: Any) -> int | None:
    if raw is None or raw == "":
        return None
    try:
        return int(raw)
    except (TypeError, ValueError):
        return None


def _parse_yes(value: Any) -> bool:
    if value is None:
        return False
    return str(value).strip().lower() in {"yes", "y", "true", "1", "да"}


def _parse_server_ids(raw: Any) -> list[int] | None:
    if raw is None or str(raw).strip() == "":
        return None
    parts = re.split(r"[,;]", str(raw))
    ids: list[int] = []
    for part in parts:
        part = part.strip()
        if not part:
            continue
        ids.append(int(part))
    return ids or None


def _parse_attachments(raw: Any) -> list[str]:
    if raw is None or str(raw).strip() == "":
        return []
    text = str(raw).strip()
    parts = re.split(r"[,;]", text)
    return [p.strip() for p in parts if p.strip()]


def _score_classname(name: str) -> int:
    """Prefer PascalCase canonical names when building lookup."""
    if not name:
        return 0
    upper = sum(1 for c in name if c.isupper())
    return upper * 10 + len(name)


def build_classname_lookup(market_dir: Path) -> dict[str, str]:
    lookup: dict[str, str] = dict(
        (k.lower(), v) for k, v in KNOWN_CLASS_ALIASES.items()
    )
    if not market_dir.is_dir():
        return lookup
    for path in sorted(market_dir.glob("*.json")):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        _collect_classnames(data, lookup)
    return lookup


def _collect_classnames(node: Any, lookup: dict[str, str]) -> None:
    if isinstance(node, dict):
        if "ClassName" in node and isinstance(node["ClassName"], str):
            name = node["ClassName"].strip()
            if name:
                key = name.lower()
                if key not in lookup or _score_classname(name) > _score_classname(
                    lookup[key]
                ):
                    lookup[key] = name
        for value in node.values():
            _collect_classnames(value, lookup)
    elif isinstance(node, list):
        for item in node:
            if isinstance(item, str) and item.strip():
                name = item.strip()
                key = name.lower()
                if key not in lookup:
                    lookup[key] = name
            else:
                _collect_classnames(item, lookup)


def normalize_classname(raw: str, lookup: dict[str, str]) -> str:
    key = raw.strip().lower()
    if key in lookup:
        return lookup[key]
    segments = key.split("_")
    normalized: list[str] = []
    for seg in segments:
        if not seg:
            continue
        if seg.isdigit():
            normalized.append(seg)
        elif len(seg) <= 3 and seg.isalpha():
            normalized.append(seg.upper())
        else:
            normalized.append(seg[0].upper() + seg[1:])
    return "_".join(normalized)


def find_market_spawn_attachments(
    vehicle_class: str, market_dir: Path
) -> list[str] | None:
    if not market_dir.is_dir():
        return None
    target = vehicle_class.lower()
    for path in sorted(market_dir.glob("*.json")):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        for item in data.get("Items") or []:
            cname = str(item.get("ClassName") or "").strip()
            if cname.lower() != target:
                continue
            attachments = item.get("SpawnAttachments")
            if isinstance(attachments, list) and attachments:
                return [str(a) for a in attachments]
    return None


def load_items_sheet(wb) -> dict[int, list[dict[str, Any]]]:
    ws = wb["items"]
    grouped: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for row in ws.iter_rows(min_row=2, values_only=True):
        offer_id = _parse_offer_id(row[0])
        if offer_id is None or offer_id in SKIP_OFFER_IDS:
            continue
        item_class = row[1]
        if item_class is None or str(item_class).strip() == "":
            continue
        qty = int(row[2]) if row[2] not in (None, "") else 1
        entry: dict[str, Any] = {
            "class": str(item_class).strip(),
            "qty": qty,
        }
        attachments = _parse_attachments(row[3])
        if attachments:
            entry["attachments"] = attachments
        grouped[offer_id].append(entry)
    return grouped


def load_offers_sheet(
    wb,
    items_by_offer: dict[int, list[dict[str, Any]]],
    market_dir: Path,
    lookup: dict[str, str],
    report: ImportReport,
    vehicle_key_classes: set[str],
) -> dict[int, dict[str, Any]]:
    ws = wb["offers"]
    offers: dict[int, dict[str, Any]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        offer_id = _parse_offer_id(row[0])
        if offer_id is None or offer_id in SKIP_OFFER_IDS:
            if offer_id in SKIP_OFFER_IDS:
                report.skipped_template.append(offer_id)
            continue

        name = str(row[1]).strip() if row[1] is not None else f"Offer {offer_id}"
        delivery_online = row[2]
        server_ids = _parse_server_ids(row[3])
        excel_type = str(row[4] or "container").strip().lower()
        container_or_class = (
            str(row[5]).strip() if row[5] is not None else ""
        )
        give_key_raw = row[6]
        spawn_distance = row[7]
        spawn_slope = row[8]
        spawn_radius = row[9]
        spawn_clearance = row[10]

        if not _parse_yes(delivery_online) and delivery_online not in (None, ""):
            pass
        elif delivery_online in (None, ""):
            report.warnings.append(
                f"offer {offer_id}: delivery_online пусто — в ЛК wargm нужен тип online/api"
            )

        if offer_id in VEHICLE_OVERRIDES:
            vehicle_class = VEHICLE_OVERRIDES[offer_id]
            if _parse_yes(give_key_raw) or give_key_raw in (None, ""):
                vehicle_key_classes.add(vehicle_class)
            offer: dict[str, Any] = {
                "type": "vehicle",
                "name": name,
                "class": vehicle_class,
            }
            if server_ids:
                offer["shop_server_ids"] = server_ids
            spawn = dict(DEFAULT_SPAWN)
            if spawn_distance not in (None, ""):
                spawn["distance_m"] = float(spawn_distance)
            if spawn_slope not in (None, ""):
                spawn["max_slope_deg"] = float(spawn_slope)
            if spawn_radius not in (None, ""):
                spawn["check_radius_m"] = float(spawn_radius)
            if spawn_clearance not in (None, ""):
                spawn["min_clearance_above_m"] = float(spawn_clearance)
            offer["spawn"] = spawn
            offers[offer_id] = offer
            continue

        dtype = excel_type if excel_type in {"container", "vehicle"} else "container"
        if dtype == "vehicle":
            vehicle_class = container_or_class
            if _parse_yes(give_key_raw):
                vehicle_key_classes.add(str(vehicle_class))
            offer = {
                "type": "vehicle",
                "name": name,
                "class": vehicle_class,
            }
            if server_ids:
                offer["shop_server_ids"] = server_ids
            spawn = dict(DEFAULT_SPAWN)
            if spawn_distance not in (None, ""):
                spawn["distance_m"] = float(spawn_distance)
            if spawn_slope not in (None, ""):
                spawn["max_slope_deg"] = float(spawn_slope)
            if spawn_radius not in (None, ""):
                spawn["check_radius_m"] = float(spawn_radius)
            if spawn_clearance not in (None, ""):
                spawn["min_clearance_above_m"] = float(spawn_clearance)
            offer["spawn"] = spawn
            offers[offer_id] = offer
            continue

        container = container_or_class or "SeaChest"
        offer_items = items_by_offer.get(offer_id, [])
        if not offer_items:
            report.errors.append(f"offer {offer_id}: container без items")

        qty_sum = sum(i.get("qty", 1) for i in offer_items)
        if qty_sum > 7:
            report.warnings.append(
                f"offer {offer_id}: {qty_sum} предметов в {container} — риск переполнения cargo"
            )

        offer = {
            "type": "container",
            "name": name,
            "container": container,
            "items": offer_items,
        }
        if server_ids:
            offer["shop_server_ids"] = server_ids
        offers[offer_id] = offer

    return offers


def build_vehicle_profiles(
    offers: dict[int, dict[str, Any]],
    market_dir: Path,
    lookup: dict[str, str],
    report: ImportReport,
    vehicle_key_classes: set[str],
) -> dict[str, dict[str, Any]]:
    profiles: dict[str, dict[str, Any]] = {}
    vehicle_classes: set[str] = set()
    for offer in offers.values():
        if offer.get("type") == "vehicle":
            cls = offer.get("class")
            if cls:
                vehicle_classes.add(str(cls))

    for vehicle_class in sorted(vehicle_classes):
        raw_attachments = find_market_spawn_attachments(vehicle_class, market_dir)
        if not raw_attachments:
            report.errors.append(
                f"vehicle {vehicle_class}: SpawnAttachments не найдены в Market"
            )
            continue
        normalized = [
            normalize_classname(a, lookup) for a in raw_attachments
        ]
        for raw, norm in zip(raw_attachments, normalized):
            if raw.lower() != norm.lower():
                report.warnings.append(
                    f"vehicle {vehicle_class}: нормализация {raw!r} → {norm!r}"
                )
        profiles[vehicle_class] = {
            "spawn_attachments": normalized,
            "fluids": {"fuel": 1.0, "coolant": 1.0},
        }
        if vehicle_class in vehicle_key_classes:
            profiles[vehicle_class]["cargo_items"] = [
                {"class": "ExpansionCarKey", "qty": 1}
            ]
    return profiles


class ImportReport:
    def __init__(self) -> None:
        self.warnings: list[str] = []
        self.errors: list[str] = []
        self.skipped_template: list[int] = []

    def print_summary(self, offer_count: int, profile_count: int) -> None:
        if self.skipped_template:
            print(
                f"Skipped template offers: {sorted(self.skipped_template)}"
            )
        print(f"Imported {offer_count} offers, {profile_count} vehicle profiles")
        for w in self.warnings:
            print(f"WARNING: {w}")
        for e in self.errors:
            print(f"ERROR: {e}")


def write_offers_yaml(path: Path, offers: dict[int, dict[str, Any]]) -> None:
    header = (
        "# Каталог товаров wargm → выдача в DayZ\n"
        "# Сгенерировано: catalog/import/import-offers-from-xlsx.py\n"
        "# Не редактировать вручную — править Excel и перезапускать импорт.\n"
    )
    data = {"offers": {oid: offers[oid] for oid in sorted(offers)}}
    body = yaml.dump(
        data,
        allow_unicode=True,
        default_flow_style=False,
        sort_keys=False,
        width=120,
    )
    path.write_text(header + "\n" + body, encoding="utf-8")


def write_profiles_yaml(path: Path, profiles: dict[str, dict[str, Any]]) -> None:
    header = (
        "# Профили комплектации техники — SpawnAttachments из Expansion Market.\n"
        "# Сгенерировано: catalog/import/import-offers-from-xlsx.py\n"
    )
    data = {"profiles": profiles}
    body = yaml.dump(
        data,
        allow_unicode=True,
        default_flow_style=False,
        sort_keys=False,
        width=120,
    )
    path.write_text(header + "\n" + body, encoding="utf-8")


def run_import(
    input_path: Path,
    offers_out: Path,
    profiles_out: Path,
    market_dir: Path,
) -> ImportReport:
    report = ImportReport()
    if not input_path.is_file():
        report.errors.append(f"input not found: {input_path}")
        return report

    lookup = build_classname_lookup(market_dir)
    wb = load_workbook(input_path, read_only=True, data_only=True)
    vehicle_key_classes: set[str] = set()
    try:
        items_by_offer = load_items_sheet(wb)
        offers = load_offers_sheet(
            wb, items_by_offer, market_dir, lookup, report, vehicle_key_classes
        )
    finally:
        wb.close()

    profiles = build_vehicle_profiles(
        offers, market_dir, lookup, report, vehicle_key_classes
    )

    if not report.errors:
        write_offers_yaml(offers_out, offers)
        write_profiles_yaml(profiles_out, profiles)

    report.print_summary(len(offers), len(profiles))
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="Import wargm offers from Excel")
    parser.add_argument(
        "--input",
        type=Path,
        default=DEFAULT_INPUT,
        help="Path to offers-template.xlsx",
    )
    parser.add_argument(
        "--offers-out",
        type=Path,
        default=CATALOG_DIR / "offers.yaml",
    )
    parser.add_argument(
        "--profiles-out",
        type=Path,
        default=CATALOG_DIR / "vehicle_profiles.yaml",
    )
    parser.add_argument(
        "--market-dir",
        type=Path,
        default=DEFAULT_MARKET_DIR,
        help="Expansion Market JSON directory",
    )
    args = parser.parse_args()

    report = run_import(
        args.input,
        args.offers_out,
        args.profiles_out,
        args.market_dir,
    )
    return 1 if report.errors else 0


if __name__ == "__main__":
    sys.exit(main())
