# -*- coding: utf-8 -*-
"""One-off: generate offers-template.xlsx. Run: python build-offers-template.py"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent / "offers-template.xlsx"


def style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for col in range(1, sheet.max_column + 1):
        letter = get_column_letter(col)
        max_len = 12
        for row in sheet.iter_rows(
            min_row=1, max_row=sheet.max_row, min_col=col, max_col=col
        ):
            val = row[0].value
            if val is not None:
                max_len = max(max_len, min(len(str(val)) + 2, 60))
        sheet.column_dimensions[letter].width = max_len
    sheet.freeze_panes = "A2"


def main() -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "offers"
    ws.append(
        [
            "offer_id",
            "name_wargm",
            "delivery_online",
            "shop_server_ids",
            "type",
            "container_or_vehicle_class",
            "give_key",
            "spawn_distance_m",
            "spawn_max_slope_deg",
            "spawn_check_radius_m",
            "spawn_min_clearance_m",
            "notes",
        ]
    )
    for row in [
        [
            264775,
            "test",
            "yes",
            "",
            "container",
            "SeaChest",
            "",
            "",
            "",
            "",
            "",
            "TaloonBag_Blue — см. лист items",
        ],
        [
            264776,
            "test2",
            "yes",
            "10001",
            "vehicle",
            "OffroadHatchback",
            "yes",
            6,
            15,
            2.5,
            "",
            "wargm delivery=api; комплектация в vehicle_profiles",
        ],
        [
            12346,
            "5000 рублей (пример)",
            "yes",
            "",
            "container",
            "SeaChest",
            "",
            "",
            "",
            "",
            "",
            "Money_Ruble5000 x1",
        ],
        [
            5020,
            "Тест: бочка",
            "yes",
            "",
            "container",
            "Barrel_Green",
            "",
            "",
            "",
            "",
            "",
            "mock dev; BandageDressing x3",
        ],
        [
            5021,
            "Тест: ящик FL camo",
            "yes",
            "",
            "container",
            "LA_SeaChest_camo",
            "",
            "",
            "",
            "",
            "",
            "mock dev; demo; BandageDressing x5",
        ],
        [
            12351,
            "UH-1H (пример)",
            "yes",
            "10001",
            "vehicle",
            "RFFSHeli_UH1H",
            "yes",
            12,
            10,
            5,
            8,
            "только Chernarus",
        ],
    ]:
        ws.append(row)

    ws2 = wb.create_sheet("items")
    ws2.append(["offer_id", "item_class", "qty", "attachments"])
    for row in [
        [264775, "TaloonBag_Blue", 1, ""],
        [12346, "Money_Ruble5000", 1, ""],
        [5020, "BandageDressing", 3, ""],
        [5021, "BandageDressing", 5, ""],
    ]:
        ws2.append(row)

    ws3 = wb.create_sheet("vehicle_profiles")
    ws3.append(
        ["vehicle_class", "spawn_attachments", "fuel", "coolant", "notes"]
    )
    ws3.append(
        [
            "OffroadHatchback",
            "HatchbackWheel,HatchbackWheel,HatchbackWheel,HatchbackWheel,HatchbackWheel,"
            "CarBattery,CarRadiator,SparkPlug,HatchbackHood,HatchbackTrunk,"
            "HatchbackDoors_Driver,HatchbackDoors_CoDriver,HeadlightH7,HeadlightH7",
            1.0,
            1.0,
            "копия Expansion Market Cars.json",
        ]
    )

    ws4 = wb.create_sheet("справка")
    for row in [
        ["Поле", "Источник", "Пояснение"],
        ["offer_id", "wargm ЛК", "ID предложения в магазине — ключ в offers.yaml"],
        ["name_wargm", "wargm ЛК", "Название для поля name в каталоге"],
        [
            "delivery_online",
            "wargm ЛК",
            "Должно быть yes; в ЛК тип доставки = online",
        ],
        [
            "shop_server_ids",
            "wargm ЛК",
            "Пусто = все сервера; иначе 10001,10002,10003 через запятую",
        ],
        ["type", "вы", "container или vehicle"],
        ["container_or_vehicle_class", "вы", "Container_Base: SeaChest, Barrel_Green, LA_SeaChest_camo… или ClassName техники"],
        ["give_key", "вы", "yes/no только для vehicle — пишет ExpansionCarKey в vehicle_profiles cargo_items"],
        ["spawn_*", "вы", "только vehicle; вертолёт — min_clearance_m"],
        ["", "", ""],
        ["Серверы wargm (без Alteria)", "ID", "Карта"],
        ["demo_server_alpha", 10001, "Demo server alpha"],
        ["demo_server_beta", 10002, "Demo server beta"],
        ["demo_server_gamma", 10003, "Demo server gamma"],
    ]:
        ws4.append(row)

    for sheet in [ws, ws2, ws3, ws4]:
        style_sheet(sheet)

    wb.save(OUT)
    print(f"Created: {OUT}")


if __name__ == "__main__":
    main()
